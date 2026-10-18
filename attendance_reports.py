"""Attendance and class-event report generation used by the Flask dashboard."""
from io import BytesIO
from pathlib import Path
import textwrap

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

REQUIRED = ["Batch", "Scanned_date"]
BATCH_ALIASES = {}
COUNT_ONLY_STATUS = None


def _parse_dates(column):
    values = column.astype(str).str.strip()
    iso = values.str.match(r"^\d{4}-\d{2}-\d{2}")
    dates = pd.Series(pd.NaT, index=values.index, dtype="datetime64[ns]")
    dates[iso] = pd.to_datetime(values[iso].str[:10], format="%Y-%m-%d", errors="coerce")
    dates[~iso] = pd.to_datetime(values[~iso], format="mixed", dayfirst=True, errors="coerce")
    return dates.dt.normalize()


def load(source):
    """Read the class-event Excel export, entirely in memory."""
    sheets = pd.read_excel(source, sheet_name=None, dtype=str)
    dataframe = next((sheet for sheet in sheets.values()
                      if set(REQUIRED) <= {str(column).strip() for column in sheet.columns}), None)
    if dataframe is None:
        found = list(next(iter(sheets.values())).columns) if sheets else []
        raise KeyError(f"Batch / Scanned_date (columns found: {found})")
    dataframe = dataframe.copy()
    dataframe.columns = [str(column).strip() for column in dataframe.columns]
    dataframe["Batch"] = (dataframe["Batch"].fillna("(no batch)").str.split().str.join(" ")
                          .replace("", "(no batch)").replace(BATCH_ALIASES))
    if "Admission_no" in dataframe.columns:
        dataframe["Admission_no"] = dataframe["Admission_no"].str.strip().replace("", pd.NA)
    dataframe["Date"] = _parse_dates(dataframe["Scanned_date"])
    skipped_rows = int(dataframe["Date"].isna().sum())
    dataframe = dataframe[dataframe["Date"].notna()]
    if dataframe.empty:
        raise ValueError("No valid Scanned_date values found in this file.")
    dataframe.attrs["skipped_rows"] = skipped_rows
    if "Admission_no" not in dataframe.columns:
        if "Student_dynamic_id" not in dataframe.columns:
            raise ValueError("Missing student identifier: Admission_no or Student_dynamic_id is required.")
        dataframe["Admission_no"] = dataframe["Student_dynamic_id"].str.strip().replace("", pd.NA)
    if COUNT_ONLY_STATUS:
        dataframe = dataframe[dataframe["Status"].str.strip() == COUNT_ONLY_STATUS]
    return dataframe


def summarize(dataframe):
    days = pd.date_range(dataframe["Date"].min(), dataframe["Date"].max())
    daily = dataframe.groupby(["Batch", "Date"]).size().unstack(fill_value=0).reindex(columns=days, fill_value=0)
    daily = daily.loc[daily.sum(axis=1).sort_values(ascending=False).index]
    cumulative = daily.cumsum(axis=1)
    unique_daily = dataframe.groupby("Date")["Admission_no"].nunique().reindex(days, fill_value=0)
    return {
        "days": [day.strftime("%d-%m-%Y") for day in days],
        "batches": [{"batch": batch, "daily": [int(value) for value in daily.loc[batch]],
                     "cumulative": [int(value) for value in cumulative.loc[batch]],
                     "total": int(daily.loc[batch].sum())} for batch in daily.index],
        "total_daily": [int(value) for value in daily.sum()],
        "total_cumulative": [int(value) for value in cumulative.sum()],
        "grand_total": int(daily.values.sum()),
        "unique_daily": [int(value) for value in unique_daily],
        "unique_overall": int(dataframe["Admission_no"].nunique()),
        "skipped_rows": int(dataframe.attrs.get("skipped_rows", 0)),
        "missing_ids": int(dataframe["Admission_no"].isna().sum()),
    }


def build_excel(dataframe, output, event_name):
    """Write the original Daily and Cumulative workbook layout to a buffer."""
    report = summarize(dataframe)
    days = pd.date_range(dataframe["Date"].min(), dataframe["Date"].max())
    workbook = Workbook()
    thin = Side(style="thin", color="BBBBBB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill, total_fill = PatternFill("solid", fgColor="1F3864"), PatternFill("solid", fgColor="D9E1F2")

    def font(**options): return Font(name="Arial", size=10, **options)

    def write_sheet(sheet, title, cumulative):
        sheet["A1"] = f"{event_name} - {title}"
        sheet["A1"].font = Font(name="Arial", size=13, bold=True)
        sheet["A2"] = f"Report date: {days[-1]:%d-%m-%Y}   |   Event day: {len(days)}"
        sheet["A2"].font = font(italic=True)
        headers = ["Batch"] + [f"Day {index + 1}\n{day:%d-%m-%Y}" for index, day in enumerate(days)]
        if not cumulative: headers.append("Total")
        for column, value in enumerate(headers, 1):
            cell = sheet.cell(4, column, value)
            cell.font, cell.fill, cell.border = font(bold=True, color="FFFFFF"), header_fill, border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.row_dimensions[4].height = 32
        for row_index, batch in enumerate(report["batches"], 5):
            values = batch["cumulative"] if cumulative else batch["daily"]
            sheet.cell(row_index, 1, batch["batch"])
            for column, value in enumerate(values, 2):
                if cumulative:
                    letter = get_column_letter(column)
                    sheet.cell(row_index, column, f"=SUM(Daily!$B{row_index}:{letter}{row_index})")
                else:
                    sheet.cell(row_index, column, value)
            if not cumulative:
                sheet.cell(row_index, len(days) + 2, f"=SUM(B{row_index}:{get_column_letter(len(days)+1)}{row_index})")
        total_row = 5 + len(report["batches"])
        sheet.cell(total_row, 1, "Grand Total")
        last_column = len(headers)
        for column in range(2, last_column + 1):
            letter = get_column_letter(column)
            sheet.cell(total_row, column, f"=SUM({letter}5:{letter}{total_row - 1})")
        for row in range(5, total_row + 1):
            for column in range(1, last_column + 1):
                cell = sheet.cell(row, column); cell.font = font(bold=row == total_row); cell.border = border
                if column > 1: cell.alignment = Alignment(horizontal="center")
                if row == total_row: cell.fill = total_fill
        sheet.column_dimensions["A"].width = 40
        for column in range(2, last_column + 1): sheet.column_dimensions[get_column_letter(column)].width = 14
        sheet.freeze_panes = "B5"
        return total_row

    daily_sheet = workbook.active; daily_sheet.title = "Daily"
    total_row = write_sheet(daily_sheet, "Daily Attendance (scans per day)", False)
    unique_row = total_row + 2
    daily_sheet.cell(unique_row, 1, "Unique students (distinct admission no.)").font = font(bold=True)
    for column, value in enumerate(report["unique_daily"], 2):
        daily_sheet.cell(unique_row, column, value).alignment = Alignment(horizontal="center")
    daily_sheet.cell(unique_row, len(days) + 2, report["unique_overall"]).font = font(bold=True)
    cumulative_sheet = workbook.create_sheet("Cumulative")
    write_sheet(cumulative_sheet, "Cumulative Attendance (running total)", True)
    workbook.save(output)


def build_image(dataframe, output, event_name):
    """Generate the shareable PNG report, preserving the original export feature."""
    from PIL import Image, ImageDraw, ImageFont
    report = summarize(dataframe)
    day_count = len(report["days"])
    width, margin = max(1500, 600 + (day_count + 1) * 125), 64
    batch_width = 500
    cell_width = (width - margin * 2 - batch_width) / (day_count + 1)
    font_dir = Path("C:/Windows/Fonts")
    def font(size, bold=False):
        for name in (["arialbd.ttf", "DejaVuSans-Bold.ttf"] if bold else ["arial.ttf", "DejaVuSans.ttf"]):
            try: return ImageFont.truetype(str(font_dir / name) if (font_dir / name).exists() else name, size)
            except OSError: pass
        return ImageFont.load_default()
    title_lines = textwrap.wrap(event_name, max(20, int((width - 128) / 29))) or ["Class Event"]
    row_heights = [max(58, 28 * len(textwrap.wrap(batch["batch"], 35)) + 20) for batch in report["batches"]]
    header_height = 190 + (len(title_lines) - 1) * 56
    height = header_height + 360 + 2 * (126 + sum(row_heights)) + 600
    image = Image.new("RGB", (width, height), "#eef3f8"); draw = ImageDraw.Draw(image)
    navy, teal, muted = "#132d46", "#087e8b", "#607489"
    draw.rectangle((0, 0, width, header_height), fill=navy)
    draw.text((margin, 30), "ATTENDANCE REPORT", font=font(19, True), fill="#77d5dc")
    y = 69
    for line in title_lines:
        draw.text((margin, y), line, font=font(44, True), fill="white"); y += 56
    draw.text((margin, y + 9), f"{report['days'][0]} to {report['days'][-1]}    /    Event day {day_count}", font=font(22), fill="#c7d8e8")
    y = header_height + 28
    cards = [("TOTAL SCANS", report["grand_total"]), ("UNIQUE STUDENTS", report["unique_overall"]),
             ("BATCHES", len(report["batches"])), ("LATEST DAY SCANS", report["total_daily"][-1])]
    card_width = (width - margin * 2 - 54) / 4
    for index, (label, value) in enumerate(cards):
        x = margin + index * (card_width + 18)
        draw.rounded_rectangle((x, y, x + card_width, y + 113), radius=16, fill="white")
        draw.text((x + 22, y + 18), label, font=font(16, True), fill=muted)
        draw.text((x + 22, y + 47), f"{value:,}", font=font(39, True), fill=navy)
    y += 141; draw.text((margin, y), "Daily attendance trend", font=font(24, True), fill=navy); y += 48
    peak = max(report["total_daily"], default=1)
    for index, value in enumerate(report["total_daily"]):
        x, bar_width = margin + index * ((width - 2 * margin) / day_count), max(15, min(120, (width - 2 * margin) / day_count - 25))
        bar_height = 78 * value / peak
        draw.rounded_rectangle((x, y + 85 - bar_height, x + bar_width, y + 85), radius=5, fill=teal)
        draw.text((x, y + 85 - bar_height - 26), str(value), font=font(19, True), fill=navy)
        draw.text((x, y + 93), report["days"][index][:5], font=font(16), fill=muted)
    y += 151
    def table(y_position, title, cumulative=False):
        draw.text((margin, y_position), title, font=font(27, True), fill=navy); y_position += 49
        headers = ["Batch"] + [f"Day {index + 1}\n{day[:5]}" for index, day in enumerate(report["days"])] + ["Total"]
        def row(y_row, values, row_height, fill, color=navy, bold=False):
            draw.rectangle((margin, y_row, width - margin, y_row + row_height), fill=fill)
            for index, value in enumerate(values):
                left = margin if index == 0 else margin + batch_width + (index - 1) * cell_width
                cell_width_here = batch_width if index == 0 else cell_width
                label = "\n".join(textwrap.wrap(str(value), 35)) if index == 0 else str(value)
                current_font = font(19 if index == 0 else 21, bold)
                bounds = draw.multiline_textbbox((0, 0), label, font=current_font, spacing=5)
                x = left + 18 if index == 0 else left + (cell_width_here - (bounds[2] - bounds[0])) / 2
                draw.multiline_text((x, y_row + (row_height - (bounds[3] - bounds[1])) / 2 - bounds[1]), label, font=current_font, fill=color, spacing=5)
            draw.line((margin, y_row + row_height, width - margin, y_row + row_height), fill="#dfe7ef", width=1)
            return y_row + row_height
        y_position = row(y_position, headers, 76, navy, "white", True)
        for index, batch in enumerate(report["batches"]):
            values = batch["cumulative"] if cumulative else batch["daily"]
            y_position = row(y_position, [batch["batch"]] + values + [batch["total"]], row_heights[index], "white" if index % 2 == 0 else "#f5f8fb")
        totals = report["total_cumulative"] if cumulative else report["total_daily"]
        y_position = row(y_position, ["Grand Total"] + totals + [report["grand_total"]], 60, "#dceef0", navy, True)
        if not cumulative: y_position = row(y_position, ["Unique students"] + report["unique_daily"] + [report["unique_overall"]], 58, "#e5ebf3")
        return y_position
    y = table(y, "01   Daily attendance", False) + 40
    y = table(y, "02   Cumulative attendance", True) + 28
    draw.text((margin, y), "Scans count every attendance record. Unique students count each admission number once.", font=font(18), fill=muted)
    image.crop((0, 0, width, y + 62)).save(output, format="PNG")


def excel_bytes(dataframe, event_name):
    output = BytesIO(); build_excel(dataframe, output, event_name); return output.getvalue()


def image_bytes(dataframe, event_name):
    output = BytesIO(); build_image(dataframe, output, event_name); return output.getvalue()
